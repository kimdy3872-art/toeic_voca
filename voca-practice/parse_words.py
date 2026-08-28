import os
import glob
import pandas as pd
import json
import re
import urllib.request
import urllib.parse
import openpyxl

# Built-in TOEIC Dictionary Overrides for common high-frequency TOEIC vocabulary
TOEIC_DICTIONARY = {
    # Recruitment & HR
    'applicant': '지원자, 응모자',
    'application': '지원(서), 신청',
    'apprentice': '견습생, 수습생',
    'candidate': '후보자, 지원자',
    'resume': '이력서',
    'cover letter': '자기소개서',
    'requirement': '자격 요건, 필요조건',
    'qualification': '자격, 자격증',
    'experience': '경력, 경험',
    'hire': '고용하다, 채용하다',
    'recruit': '채용하다, 모집하다',
    'payroll': '급여 명세서, 총급여',
    'salary': '급여, 봉급',
    'wage': '임금, 시급',
    'benefit': '혜택, 복리후생',
    'promotion': '승진, 홍보, 판촉',
    'relocate': '이전하다, 전근 가다',
    'resign': '사임하다, 사직하다',
    'retire': '퇴직하다, 은퇴하다',
    'severance': '퇴직금',
    
    # Management & Strategy
    'supervise': '감독하다, 관리하다',
    'supervisor': '상사, 감독관',
    'executive': '임원, 경영진',
    'directory': '주소록, 안내판',
    'efficiency': '효율성, 능률',
    'implement': '실행하다, 이행하다',
    'strategy': '전략, 계획',
    'accomplish': '성취하다, 달성하다',
    'evaluate': '평가하다',
    'performance': '성과, 실적, 공연',
    'objective': '목표, 목적',
    'delegate': '위임하다, 대표자',
    'initiative': '주도권, 솔선, 계획',
    'headquarters': '본사, 본부',
    'branch': '지사, 지점',
    'subsidiary': '자회사',
    'merge': '합병하다',
    'acquisition': '인수, 매수',
    
    # Finance & Commerce
    'revenue': '수익, 세입',
    'expenditure': '지출, 비용',
    'budget': '예산',
    'fiscal': '회계의, 재정의',
    'reimburse': '상환하다, 변제하다',
    'invoice': '청구서, 송장',
    'receipt': '영수증',
    'deposit': '보증금, 예금하다',
    'statement': '내역서, 명세서',
    'transaction': '거래, 매매',
    'audit': '회계 감사',
    'profit': '이익, 수익',
    'deficit': '적자, 부족액',
    'surplus': '흑자, 여유',
    'installment': '할부, 분입금',
    'balance': '잔액, 균형',
    
    # Sales & Marketing
    'survey': '설문조사, 조사하다',
    'analysis': '분석, 검토',
    'respondent': '응답자',
    'monopoly': '(상품의) 독점, 전매',
    'competition': '경쟁, 시합',
    'consistently': '항상, 일관되게',
    'demand': '수요, 요구하다',
    'do one\'s utmost': '전력을 다하다',
    'expand': '확장하다, 넓히다',
    'advanced': '고급의, 진보한, 앞선',
    'postpone': '연기하다, 미루다',
    'additional': '추가의, 부가의',
    'appreciate': '고맙게 생각하다, 감상하다',
    'demonstration': '설명, 시연, 시위',
    'examine': '조사하다, 검사하다',
    'effective': '효과적인, 발효되는',
    'closely': '면밀히, 엄밀히',
    'reserve': '예약하다, 지정하다',
    'cooperate': '협력하다',
    'consecutive': '연속적인, 계속되는',
    'expectation': '예상, 기대',
    'publicize': '공표하다, 홍보하다',
    'raise': '높이다, 올리다, 제기하다',
    'affect': '영향을 미치다',
    'target': '목표, 겨냥하다',
    'campaign': '운동, 캠페인',
    'probable': '개연성이 높은, 유망한',
    'focus': '집중시키다, 집중하다',
    'seasonal': '계절의, 계절적인',
    'impact': '영향, 충격',
    'comparison': '비교',
    'gap': '격차, 공백',
    'mounting': '증가하는, 오르는',
    'reflective': '반영하는, 반사하는',
    'compliance': '(규정·법률의) 준수',
    'eligible': '자격이 있는'
}

def has_korean(text):
    if not text:
        return False
    return bool(re.search(r'[\uac00-\ud7a3]', str(text)))

def is_english_word(text):
    if not text:
        return False
    s = str(text).strip()
    if not s or s.lower() == 'nan':
        return False
    if has_korean(s):
        return False
    return bool(re.search(r'[a-zA-Z]', s))

def get_toeic_korean_meaning(word):
    clean = word.strip().lower()
    if clean in TOEIC_DICTIONARY:
        return TOEIC_DICTIONARY[clean]
    
    url = f'https://translate.googleapis.com/translate_a/single?client=gtx&sl=en&tl=ko&dt=t&q={urllib.parse.quote(word)}'
    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
    try:
        with urllib.request.urlopen(req) as resp:
            data = json.loads(resp.read().decode('utf-8'))
            return data[0][0][0]
    except Exception:
        return '뜻 자동생성 실패'

def auto_update_excel_file(file_path, updated_sheets_data):
    """
    If missing Korean meanings were generated, auto-populate the Excel spreadsheet
    so the user gets their Excel file updated with the Korean meanings as well.
    """
    try:
        wb = openpyxl.load_workbook(file_path)
        modified = False
        
        for sheet_name, categories in updated_sheets_data.items():
            if sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                col_indices = {}
                for col_idx in range(1, ws.max_column + 1):
                    val = str(ws.cell(row=1, column=col_idx).value or '').strip()
                    if val:
                        col_indices[val] = col_idx
                
                for cat_name, word_list in categories.items():
                    if cat_name in col_indices:
                        col_idx = col_indices[cat_name]
                        current_row = 2
                        for item in word_list:
                            ws.cell(row=current_row, column=col_idx, value=item['english'])
                            ws.cell(row=current_row + 1, column=col_idx, value=item['korean'])
                            current_row += 2
                        modified = True
                        
        if modified:
            wb.save(file_path)
            print(f"  [Excel Auto-Fill] Updated '{os.path.basename(file_path)}' with generated Korean meanings!")
    except Exception as e:
        print(f"  [Excel Auto-Fill Notice] Could not update Excel file: {e}")

def parse_excel_files():
    # Resolved from this file, not hardcoded — see the note in run.py.
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    web_dir = os.path.join(base_dir, "voca-practice", "web")
    os.makedirs(web_dir, exist_ok=True)
    
    excel_files = glob.glob(os.path.join(base_dir, "*.xlsx"))
    
    parsed_data = {}
    
    for file_path in excel_files:
        file_name = os.path.basename(file_path)
        # Skip temporary files created by Excel
        if file_name.startswith("~$"):
            continue
            
        print(f"Parsing file: {file_name}")
        updated_sheet_data = {}
        file_needed_auto_translation = False
        
        try:
            excel_file = pd.ExcelFile(file_path)
            
            for sheet_name in excel_file.sheet_names:
                if sheet_name == "오답노트":
                    continue
                    
                df = pd.read_excel(file_path, sheet_name=sheet_name)
                updated_sheet_data[sheet_name] = {}
                
                for col in df.columns:
                    col_str = str(col).strip()
                    if not col_str or col_str.startswith("Unnamed:"):
                        continue
                        
                    values = [col] + df[col].dropna().tolist()
                    values = [str(v).strip() for v in values if str(v).strip() and str(v).strip().lower() != 'nan']
                    
                    if len(values) < 2:
                        continue
                        
                    category = col_str
                    words = []
                    
                    i = 1
                    while i < len(values):
                        curr = values[i]
                        nxt = values[i+1] if i + 1 < len(values) else None
                        
                        if is_english_word(curr):
                            if nxt and has_korean(nxt):
                                words.append({
                                    "english": curr,
                                    "korean": nxt
                                })
                                i += 2
                            else:
                                kor_meaning = get_toeic_korean_meaning(curr)
                                print(f"  [Auto Meaning] '{curr}' ➔ '{kor_meaning}'")
                                words.append({
                                    "english": curr,
                                    "korean": kor_meaning
                                })
                                file_needed_auto_translation = True
                                i += 1
                        elif has_korean(curr):
                            # Skip unpaired Korean text
                            i += 1
                        else:
                            i += 1
                    
                    if words:
                        if category in parsed_data:
                            parsed_data[category].extend(words)
                        else:
                            parsed_data[category] = words
                            
                        updated_sheet_data[sheet_name][category] = words
                        print(f"  Category '{category}': Loaded {len(words)} word pairs")
            
            if file_needed_auto_translation:
                auto_update_excel_file(file_path, updated_sheet_data)
                
        except Exception as e:
            print(f"Error parsing {file_name}: {e}")
            
    # Load incorrect words from SQLite DB
    incorrect_words = []
    db_path = os.path.join(base_dir, "voca-practice", "voca.db")
    if os.path.exists(db_path):
        import sqlite3
        try:
            conn = sqlite3.connect(db_path)
            conn.row_factory = sqlite3.Row
            cursor = conn.cursor()
            cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='incorrect_words'")
            if cursor.fetchone():
                # Skip tombstoned rows (words the user has since answered correctly).
                # Guarded because this script also runs standalone, before run.py's
                # init_db() has had a chance to add the column.
                cursor.execute("PRAGMA table_info(incorrect_words)")
                has_deleted = any(col[1] == "deleted" for col in cursor.fetchall())
                if has_deleted:
                    cursor.execute("""
                        SELECT english, korean, category FROM incorrect_words
                        WHERE COALESCE(deleted, 0) = 0
                    """)
                else:
                    cursor.execute("SELECT english, korean, category FROM incorrect_words")
                rows = cursor.fetchall()
                for row in rows:
                    incorrect_words.append({
                        "english": row["english"],
                        "korean": row["korean"],
                        "category": row["category"]
                    })
                print(f"  Loaded {len(incorrect_words)} words from SQLite 'voca.db'")
            conn.close()
        except Exception as e_db:
            print(f"  Error loading from SQLite DB: {e_db}")
            
    if incorrect_words:
        parsed_data["오답노트"] = incorrect_words

    output_json_path = os.path.join(web_dir, "data.json")
    with open(output_json_path, "w", encoding="utf-8") as f:
        json.dump(parsed_data, f, ensure_ascii=False, indent=2)
        
    print(f"\nSuccessfully wrote parsed vocabulary data to {output_json_path}")
    print(f"Total categories loaded: {len(parsed_data)}")

if __name__ == "__main__":
    parse_excel_files()
